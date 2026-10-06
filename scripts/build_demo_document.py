"""Build the screenshot walkthrough with the bundled artifact Python runtime."""
import io
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'demo'
STORY = json.loads((OUT / 'story.json').read_text(encoding='utf-8'))
BOUNDARIES = [
    '本轮使用 demo 规则模式与 BM25。页眉虽显示模型配置名 deepseek-flash，但四条工单的模型调用和输入输出 token 均为 0；没有真实退款。',
    '截图展示正常流程和当前监控，未在本步注入崩溃、断库或超时。此类恢复证据见第三阶段报告，不能从截图推出任意故障均能自愈。',
    '引用匹配不等于通用语义支持性已验证；人工标签复核仍待完成。reranker 经过实验未启用。文本 PDF 支持维护命令导入，不含网页上传或 OCR；本图展示 Markdown 政策，未重新演示 PDF。',
    '共享角色授权码不是个人账号或租户隔离。本机主环境仍为 step36；截图来自独立的新版本 demo，不能当作主环境已升级或生产上线。'
]
DATA = [
    ['RF-1001', '299 元 · 签收 3 天 · 未使用', '首次 refunded；再次 already_refunded'],
    ['RF-1004', '159 元 · 签收 3 天 · 已使用', 'awaiting_approval → 批准后 refunded'],
    ['RF-1002', '129 元 · 签收 12 天 · 已使用', 'auto_rejected；无退款记录'],
]


def font(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def paragraph(doc, text, label=None):
    p = doc.add_paragraph()
    if label:
        p.add_run(label + '  ').bold = True
    p.add_run(text)
    return p


def main():
    doc = Document()
    # Remove decorative paragraph borders inherited from Word's base template.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    section.header_distance = section.footer_distance = Inches(.3)
    for name, size, bold in [('Normal',11,False), ('Title',26,True), ('Heading 1',20,True),
                              ('Heading 2',14,True), ('Caption',9,False)]:
        font(doc.styles[name],size,bold)
    normal = doc.styles['Normal'].paragraph_format
    normal.space_after = Pt(9)
    normal.line_spacing = 1.2
    for name in ('Title','Heading 1','Heading 2'):
        doc.styles[name].paragraph_format.space_after = Pt(12)
    header = section.header.paragraphs[0]
    header.text = 'ResolveFlow  /  第四阶段截图演示  /  合成数据与模拟退款'
    header.style = doc.styles['Caption']
    footer = section.footer.paragraphs[0]
    footer.text = '2026-09-23    ·    4.2    ·    '
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    footer.style = doc.styles['Caption']
    doc.core_properties.title = STORY['title']
    doc.core_properties.subject = '真实页面截图与约 2 分 40 秒讲解顺序'
    doc.core_properties.author = 'ResolveFlow'
    doc.add_heading(STORY['title'],0)
    paragraph(doc,'用一组真实页面截图，展示售后诉求如何经过规则核验、人工审批和后台执行，并通过台账避免重复模拟退款。展示业务处理流程与执行结果。')
    paragraph(doc,'主线按图 1 至图 8 顺序讲解，建议约 2 分 40 秒。图 9 为可选拒绝分支对照。时长为脚本预算，未进行真人计时。')
    doc.add_heading('演示数据与结果',1)
    table=doc.add_table(rows=1, cols=3)
    for cell,text in zip(table.rows[0].cells,['合成订单','订单事实','实测处理结果']): cell.text=text
    for row in DATA:
        for cell,text in zip(table.add_row().cells,row): cell.text=text
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(6)
                p.paragraph_format.space_before=Pt(6)
                for run in p.runs: run.font.size=Pt(10)
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9DEE3');borders.append(e)
            cell._tc.get_or_add_tcPr().append(borders)
    paragraph(doc,'最终 4 条工单、5 次执行尝试、1 条审批、2 条模拟退款，共 458 元。每条工单模型调用均为 0。RF-1003 未用于本轮演示。')
    doc.add_heading('如何阅读截图',1)
    paragraph(doc,'每页包含截图、可直接讲述的解说，以及复现操作和核对点。Word 为便于阅读裁去部分导航与空白，完整原始截图保存在同目录 screenshots，页面内容未重绘或改写。')
    paragraph(doc,'截图环境为独立项目 resolveflow-qa-demo-step42，端口 8026，运行 3.9 已测镜像。代码基线 526e166；页面上的“规则演示”表明当前没有调用该模型。')
    paragraph(doc,'复现步骤见 REPLAY.md；工单与台账证据见仓库 validation/step-4.2-2026-09-23。主环境端口 8003 不参与本次演示。')
    for n,scene in enumerate(STORY['scenes'],1):
        heading=doc.add_heading(f'{n:02d}  '+scene['title'],1)
        heading.paragraph_format.page_break_before=True
        p=paragraph(doc,scene['time']+'    ·    '+scene['role']+'角色')
        p.style=doc.styles['Caption']
        image=Image.open(OUT/'screenshots'/scene['image']).crop(scene['crop'])
        buffer=io.BytesIO();image.save(buffer,format='PNG');buffer.seek(0)
        doc.add_picture(buffer,width=Inches(7.1))
        p=doc.add_paragraph('图 '+str(n)+'  '+scene['title']+'  ·  本地独立 demo 实拍',style='Caption')
        paragraph(doc,scene['say'], '讲解')
        paragraph(doc,scene['action'], '操作')
        paragraph(doc,scene['observe'], '核对')
        if scene['run_id']:
            p=paragraph(doc,'关联工单 '+scene['run_id']);p.style=doc.styles['Caption']
        if n==8:
            doc.add_heading('展示范围',2)
            paragraph(doc,'本轮没有故障注入。监控中的毫秒数来自少量合成样例，不是生产性能承诺；当前未发现告警也不代表任何时刻都不会发生故障。')
    heading=doc.add_heading('能力边界与证据入口',1)
    heading.paragraph_format.page_break_before=True
    for line in BOUNDARIES: paragraph(doc,line)
    doc.add_heading('证据如何对应截图',2)
    paragraph(doc,'before-approval.json 保存待审批时两条工单、一条退款、零审批；after-approval.json 保存原工单恢复及审批；final-demo.json 保存四条工单的最终状态、两条退款和零模型调用。verification.json 核对前后记录、文件哈希、主环境及环境停机状态。')
    paragraph(doc,'故障与期限的历史专项证据在 validation/step-3.7-2026-09-23、step-3.8-2026-09-23 和 step-3.9-2026-09-23。成功 CI 35832707932 对应代码 526e166，不代表本次新材料已有独立 CI。')
    doc.add_heading('再次演示',2)
    paragraph(doc,'只阅读本文件即可展示。需要重新操作时，按 REPLAY.md 选择全新项目名与空闲端口；不要清空旧数据库来重置退款，不要将演示导向主环境。结束后只停止演示服务，保留卷和记录。')
    doc.save(OUT/'ResolveFlow-demo.docx')
    md=['# '+STORY['title'],'', '[下载 Word 版本](ResolveFlow-demo.docx) · [翻页与讲解指南](PRESENTING.md) · [安全复现步骤](REPLAY.md) · [4.2 实跑报告](../../validation/step-4.2-2026-09-23/REPORT.md)','',
        '真实页面截图与可直接讲述的脚本。按图 1–8 阅读，建议约 2 分 40 秒；图 9 是可选对照。'+STORY['timing_note'],'',
        '独立 demo、BM25、合成订单、模拟退款。运行镜像 resolveflow:step39-ci-fix，代码基线 `526e166`。截图环境采用 demo 模式，无生成模型调用。','',
        '## 演示数据','', '|订单|合成事实|实测结果|','|---|---|---|']
    md += ['|'+'|'.join(row)+'|' for row in DATA]
    md += ['', '最终 4 条工单、5 次执行尝试、1 条审批、2 条模拟退款（共 458 元），零模型调用。所有诉求均为“申请退款”，RF-1003 未使用。', '']
    for n,scene in enumerate(STORY['scenes'],1):
        md += ['## '+str(n)+' '+scene['title'], '',scene['time']+' · '+scene['role']+'角色','',
               '!['+scene['title']+'](screenshots/'+scene['image']+')','',
               '**讲解：** '+scene['say'],'','**操作：** '+scene['action'],'','**核对：** '+scene['observe'],'']
        if scene['run_id']: md += ['工单：`'+scene['run_id']+'`','']
    md += ['## 已知边界',''] + [x+'\n' for x in BOUNDARIES]
    md += ['## 证据入口','',
           '[审批前](../../validation/step-4.2-2026-09-23/before-approval.json)、[审批后](../../validation/step-4.2-2026-09-23/after-approval.json)、[最终状态](../../validation/step-4.2-2026-09-23/final-demo.json)、[核对结果](../../validation/step-4.2-2026-09-23/verification.json)。', '',
           '图号按讲解顺序排列；实际采集时图 9 先于图 8，图 8 先于最终图 7。Word 裁切范围记录于 [story.json](story.json)，原始 PNG 保留完整浏览器视口，均未改写内容。','',
           '截图采集于 2026-09-23，2026-10-03 完成[材料复核](../../validation/step-4.3-2026-10-03/REPORT.md)。原图与业务证据保留采集日期；当前个人账号、双工作区和容量成果见[交付导航](../../DELIVERY.md)。']
    (OUT/'README.md').write_text('\n'.join(md)+'\n',encoding='utf-8')
    print(json.dumps({'docx':str(OUT/'ResolveFlow-demo.docx'),'screenshots':len(STORY['scenes'])}))


if __name__ == '__main__':
    main()
